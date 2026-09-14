#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
码值合规检查：校验本次仿真写入的数据，其编码类字段取值是否落在码表内。

码表来源（两个权威源，取并集）：
  1. fz01_ontology.ontology_enumerations
       column_refs = [{"table":..,"column":..,"form":..}]  +  items = [{"code":..,"name":..}]
  2. fz01_governance.code_value_column_form
       (table_name, column_name) -> code_name, form
       + code_value_items(code_name, item_code, item_name)

form 语义：'编码' 存码 / '名称' 存名 / '混合' 或空 存码或名。

圈定仿真范围：读 manifest 的 pks（主键清单），建临时表 JOIN，
**不使用 ID 区间**（项目铁律）。若某表基线为 0（全表皆仿真）则跳过 JOIN 加速。

用法：
  python check_code_values.py --db fz01 --tag 模拟 \
      --manifest manifest/fz01__模拟__xxxxxxxx.json \
      --out ../reports/码值合规检查.xlsx
"""
import argparse
import json
import os
import sys
from collections import defaultdict, OrderedDict

try:
    import pymysql
except ImportError:
    sys.exit("需要 pymysql：pip install pymysql")


# ---------------------------------------------------------------- 码表装载
def load_code_maps(conn_ont, conn_gov, tables, existing):
    """返回 {(table, column): {'codes':set,'names':set,'form':str,'code_names':set,'src':set}}

    existing: 实际存在的 (table, column) 集合，用于过滤通配规则展开出的无效列。
    """
    # code_name -> (codes, names)
    by_code = defaultdict(lambda: [set(), set()])

    # --- 源1：本体枚举 ---
    col_map = {}   # (t,c) -> {'code_names':set,'form':str,'src':set}
    if conn_ont:
        cur = conn_ont.cursor()
        cur.execute("SELECT code_name, cn_name, items, column_refs FROM ontology_enumerations")
        for cn, cn_name, items, refs in cur.fetchall():
            if not refs:
                continue
            try:
                rl = json.loads(refs)
            except Exception:
                continue
            try:
                il = json.loads(items) if items else []
            except Exception:
                il = []
            for it in il:
                c, n = it.get('code'), it.get('name')
                if c is not None:
                    by_code[cn][0].add(str(c).strip())
                if n is not None:
                    by_code[cn][1].add(str(n).strip())
            if il:
                # code_name 本身也可能被 code_value_items 补充；此处登记本体自带的
                pass
            for r in rl:
                t, col = r.get('table'), r.get('column')
                if not t or not col or t not in tables:
                    continue
                d = col_map.setdefault((t, col), {'code_names': set(), 'form': '', 'src': set()})
                if cn:
                    d['code_names'].add(cn)
                if r.get('form'):
                    d['form'] = r['form']
                d['src'].add('ontology')

    # --- 源2：治理码表映射 ---
    if conn_gov:
        cur = conn_gov.cursor()
        cur.execute("SELECT code_name, item_code, item_name FROM code_value_items")
        for cn, ic, inm in cur.fetchall():
            if cn is None:
                continue
            if ic is not None:
                by_code[cn][0].add(str(ic).strip())
            if inm is not None:
                by_code[cn][1].add(str(inm).strip())
        cur.execute("SELECT table_name, column_name, code_name, form FROM code_value_column_form")
        for t, col, cn, form in cur.fetchall():
            if not col:
                continue
            # table_name='ANY' 是通配规则：按列名匹配所有表
            wild = (t == 'ANY')
            # column_name 可能是复合写法 'county_code;county_name'，拆分后逐个匹配
            names = [x.strip() for x in str(col).split(';') if x.strip()]
            targets = (sorted(tables) if wild else ([t] if t in tables else []))
            for tn in targets:
                for nc in names:
                    d = col_map.setdefault((tn, nc), {'code_names': set(), 'form': '', 'src': set()})
                    if cn:
                        d['code_names'].add(cn)
                    if form and not d['form']:
                        d['form'] = form
                    d['src'].add('code_value_column_form' + ('(ANY)' if wild else ''))

    # --- 汇总 ---
    out = {}
    for (t, c), d in col_map.items():
        if (t, c) not in existing:
            continue      # 通配规则展开出的、本表并不存在的列
        codes, names = set(), set()
        for cn in d['code_names']:
            pc, pn = by_code.get(cn, (set(), set()))
            codes |= pc
            names |= pn
        if not codes and not names:
            continue      # 无码项，无法校验
        out[(t, c)] = {
            'codes': codes, 'names': names, 'form': d['form'] or '',
            'code_names': sorted(d['code_names']), 'src': sorted(d['src']),
        }
    return out


# ---------------------------------------------------------------- 单列校验
DROP_CHARS = str.maketrans({'　': ' ', '\u3000': ' '})


# 真实库里存有 mysqldump 遗留的字面量空值哨兵，须按空处理
NULL_SENTINELS = {'\\N', 'NULL', 'null', 'None'}


def norm(v):
    if v is None:
        return None
    if isinstance(v, bytes):
        try:
            v = v.decode('utf-8')
        except Exception:
            v = str(v)
    if isinstance(v, float) and v == int(v):
        v = int(v)
    v = str(v).translate(DROP_CHARS).strip()
    if v in NULL_SENTINELS:
        return None
    return v


def check_column(cur, table, column, spec, all_sim):
    """返回 dict：取值分布与违规明细。

    用 EXISTS 半连接而非 JOIN：pk 非唯一（如日电量表按表号圈定）时 JOIN 会放大行数，
    EXISTS 既不去重也不放大，配合 _pk_.k 主键索引可走半连接。
    """
    codes, names, form = spec['codes'], spec['names'], spec['form']
    col = f"CAST(`{column}` AS CHAR)"
    if all_sim:
        q = f"SELECT {col} v, COUNT(*) n FROM `{table}` GROUP BY v"
    else:
        pk, num = spec['pk'], spec.get('pk_num')
        # 主表主键列与临时表的 collation 可能不同（库内混用 general_ci / 0900_ai_ci），
        # 直接比较会报 1267 Illegal mix of collations，故统一显式 COLLATE。
        coll = spec.get('pk_coll') or 'utf8mb4_general_ci'
        cond = (f"t.`{pk}` = p.k" if num
                else f"CAST(t.`{pk}` AS CHAR) COLLATE {coll} = p.k")
        q = (f"SELECT {col} v, COUNT(*) n FROM `{table}` t "
             f"WHERE EXISTS(SELECT 1 FROM `_pk_` p WHERE {cond}) GROUP BY v")
    cur.execute(q)
    rows = cur.fetchall()

    total = 0
    nulls = 0
    distinct_vals = 0
    bad_code = []      # 不在码里
    bad_name = []      # 不在名里
    ok_set = codes | names
    bad_any = []       # 码/名都不在
    for v, n in rows:
        nv = norm(v)
        total += n
        if nv is None or nv == '':
            nulls += n
            continue
        distinct_vals += 1
        if nv not in codes:
            bad_code.append((nv, n))
        if nv not in names:
            bad_name.append((nv, n))
        if nv not in ok_set:
            bad_any.append((nv, n))

    # 依 form 判定
    if form == '编码':
        bad = bad_code
    elif form == '名称':
        bad = bad_name
    else:
        bad = bad_any

    bad.sort(key=lambda x: -x[1])
    return {
        'table': table, 'column': column, 'form': form or '(未标注)',
        'code_names': ','.join(spec['code_names'][:3]) + ('等' if len(spec['code_names']) > 3 else ''),
        'src': '+'.join(spec['src']),
        'total': total, 'nulls': nulls, 'n_distinct': distinct_vals,
        'n_code_items': len(codes), 'n_name_items': len(names),
        'n_bad': len(bad), 'n_bad_rows': sum(n for _, n in bad),
        'bad_values': bad[:300],
        'codes_sample': sorted(codes)[:10],
        'names_sample': sorted(names)[:6],
        '_codes': sorted(codes), '_names': sorted(names),
    }


# ---------------------------------------------------------------- 违规甄别
def strip_zeros(s):
    s = str(s).strip()
    return s.lstrip('0') or '0'


def real_values(cur, table, column, spec, pks, all_sim, bad=None):
    """取真实行（非本次仿真主键）在该字段的去重取值，用于甄别。

    先整体 DISTINCT 取样；对仍不在样本内的仿真越码值，再逐值反查真实库是否存在
    ——避免高基数字段被 LIMIT 截断，把"真实同值"误判成 A 级语义错误。
    """
    if all_sim:
        return None                      # 全表皆仿真，无真实行可对照
    pk, num = spec['pk'], spec.get('pk_num')
    coll = spec.get('pk_coll') or 'utf8mb4_general_ci'
    cond = (f"t.`{pk}` = p.k" if num
            else f"CAST(t.`{pk}` AS CHAR) COLLATE {coll} = p.k")
    q = (f"SELECT DISTINCT CAST(t.`{column}` AS CHAR) FROM `{table}` t "
         f"WHERE NOT EXISTS(SELECT 1 FROM `_pk_` p WHERE {cond}) LIMIT 3000")
    try:
        cur.execute(q)
        vs = {norm(x[0]) for x in cur.fetchall()}
    except Exception:
        return None
    vs.discard(None)
    vs.discard('')
    # 逐值反查补漏
    for v in (bad or []):
        if v in vs:
            continue
        q2 = (f"SELECT 1 FROM `{table}` t "
              f"WHERE CAST(t.`{column}` AS CHAR) COLLATE {coll} = %s "
              f"AND NOT EXISTS(SELECT 1 FROM `_pk_` p WHERE {cond}) LIMIT 1")
        try:
            cur.execute(q2, (str(v),))
            if cur.fetchone():
                vs.add(v)
        except Exception:
            pass
    return vs


def classify(r, real):
    """给不合规字段定级：
    A 语义错误  —— 仿真值不在码表，真实数据也没有（真实该字段有值）
    B 格式差异  —— 去掉前导零后能对上码表码项，且真实数据不同形
    C 码表不全  —— 真实数据同样存在这些值（含同形态的短码），属码表登记不全
    D 无对照    —— 真实该字段全空，只能依据码表判定
    E 部分越码  —— 部分取值真实也有、部分独有
    """
    codes = r['_codes']
    names = r['_names']
    bad = [v for v, n in (r.get('bad_values') or [])]
    codes_nz = {strip_zeros(c) for c in codes}
    names_nz = {strip_zeros(c) for c in names}
    fmt = [v for v in bad if v not in codes and v not in names
           and (strip_zeros(v) in codes_nz or strip_zeros(v) in names_nz)]
    rest = [v for v in bad if v not in fmt]
    cs = sorted(codes)[:8]

    def fmt_is_table_issue():
        """纯格式差异时，看真实数据是否也用同一短形态 —— 是则为码表登记口径问题。"""
        return bool(real) and all(v in real for v in fmt) and not rest

    if real is None:
        return ('D 无对照(全表仿真)', fmt, rest, cs)
    if not real:
        if fmt and not rest:
            return ('B 格式差异', fmt, rest, cs)
        return ('D 无对照(真实为空)', fmt, rest, cs)
    if fmt_is_table_issue():
        return ('C 码表不全(真实同值)', fmt, [], cs)
    if not rest:
        return ('B 格式差异', fmt, [], cs)
    hit = [v for v in rest if v in real]
    miss = [v for v in rest if v not in real]
    if miss and not hit:
        return ('A 语义错误', fmt, miss, cs)
    if hit and not miss:
        return ('C 码表不全(真实同值)', fmt, hit, cs)
    return ('E 部分越码', fmt, rest, cs)


def main():
    ap = argparse.ArgumentParser(description='仿真数据码值合规检查')
    ap.add_argument('--db', default='fz01')
    ap.add_argument('--tag', default='模拟')
    ap.add_argument('--manifest', required=True, help='台账 json 路径')
    ap.add_argument('--ont-db', default='fz01_ontology')
    ap.add_argument('--gov-db', default='fz01_governance')
    ap.add_argument('--out', default=None, help='输出 xlsx（可选）')
    ap.add_argument('--json-out', default=None, help='输出明细 json（可选）')
    ap.add_argument('--limit-cols', type=int, default=0, help='调试：只查前 N 列')
    args = ap.parse_args()

    import re
    env = {}
    _envs = ['.env',
             os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', '..', '..', '.env'),
             r'D:\codex\deshu5\.env']       # deshu5 工程根目录（staging 配置）
    for p in _envs:
        if os.path.exists(p):
            for line in open(p, encoding='utf-8'):
                m = re.match(r'\s*([A-Z_]+)\s*=\s*(.*)', line)
                if m:
                    env[m.group(1)] = m.group(2).strip().strip('"').strip("'")
    pw = os.environ.get('MYSQL_PASSWORD') or env.get('MYSQL_PASSWORD')
    if not pw:
        sys.exit('未找到 MYSQL_PASSWORD（环境变量或 .env）')

    host = os.environ.get('MYSQL_HOST', '127.0.0.1')
    port = int(os.environ.get('MYSQL_PORT', '3306'))
    user = os.environ.get('MYSQL_USER', 'root')

    man = json.load(open(args.manifest, encoding='utf-8'))
    tabs = man['tables']
    tables = set(tabs)
    print(f"台账：{os.path.basename(args.manifest)}  tag={man.get('tag')}  表数={len(tables)}  "
          f"行数={sum(v['rows'] for v in tabs.values()):,}")

    # 基线快照（判断哪些表全表皆仿真）
    base = {}
    for bp in ('/tmp/simtest/base.json', 'base.json'):
        if os.path.exists(bp):
            try:
                base = json.load(open(bp, encoding='utf-8'))
                break
            except Exception:
                pass

    conn = pymysql.connect(host=host, port=port, user=user, password=pw,
                           database=args.db, charset='utf8mb4')
    try:
        conn_ont = pymysql.connect(host=host, port=port, user=user, password=pw,
                                   database=args.ont_db, charset='utf8mb4')
    except Exception:
        conn_ont = None
    try:
        conn_gov = pymysql.connect(host=host, port=port, user=user, password=pw,
                                   database=args.gov_db, charset='utf8mb4')
    except Exception:
        conn_gov = None

    # 字段类型 / 排序规则 / 字段存在性（一次性取回，供通配过滤与 _pk_ 建表使用）
    cur = conn.cursor()
    cur.execute("""SELECT table_name, column_name, data_type, collation_name
                   FROM information_schema.columns
                   WHERE table_schema=%s AND table_name IN (%s)"""
                % ('%s', ','.join(['%s'] * len(tables))),
                (args.db, *sorted(tables)))
    coltype = {}
    collmap = {}
    existing = set()
    for t, c, dt, cl in cur.fetchall():
        coltype[(t, c)] = dt
        collmap[(t, c)] = cl
        existing.add((t, c))
    NUMT = {'int', 'bigint', 'smallint', 'mediumint', 'tinyint', 'decimal', 'numeric'}

    spec_map = load_code_maps(conn_ont, conn_gov, tables, existing)
    print(f"可校验字段（有码项）：{len(spec_map)} 个，覆盖 {len({t for t, _ in spec_map})} 张表")

    work = defaultdict(list)
    for (t, c), sp in spec_map.items():
        sp = dict(sp)
        pk = tabs[t]['pk'] if isinstance(tabs[t]['pk'], str) else 'id'
        sp['pk'] = pk
        sp['pk_num'] = coltype.get((t, pk), 'varchar') in NUMT
        sp['pk_coll'] = collmap.get((t, pk)) or 'utf8mb4_general_ci'
        work[t].append((c, sp))
    if args.limit_cols:
        n = 0
        trimmed = defaultdict(list)
        for t, lst in work.items():
            for c, sp in lst:
                if n >= args.limit_cols:
                    break
                trimmed[t].append((c, sp))
                n += 1
            if n >= args.limit_cols:
                break
        work = trimmed

    results = []
    for i, (t, cols) in enumerate(sorted(work.items()), 1):
        nrows = tabs[t]['rows']
        all_sim = (base.get(t, 0) == 0) if base else False
        pks = tabs[t].get('pks') or []
        cur.execute("DROP TEMPORARY TABLE IF EXISTS `_pk_`")
        if not all_sim:
            pk = cols[0][1]['pk']
            num = cols[0][1]['pk_num']
            coll = cols[0][1]['pk_coll']
            cur.execute("CREATE TEMPORARY TABLE `_pk_` (k %s %s PRIMARY KEY)"
                        % ('BIGINT' if num else 'VARCHAR(96)',
                           '' if num else 'COLLATE ' + coll))
            CH = 5000
            for s in range(0, len(pks), CH):
                chunk = pks[s:s + CH]
                cur.executemany("INSERT IGNORE INTO `_pk_` (k) VALUES (%s)",
                                [(x if num else str(x),) for x in chunk])
        seen = set()
        for c, sp in cols:
            if c in seen:
                continue
            seen.add(c)
            try:
                r = check_column(cur, t, c, sp, all_sim)
                r['all_sim'] = all_sim
                r['pk'] = sp['pk']
                if r['n_bad'] > 0:
                    # 与真实数据对照，判定该违规是仿真缺陷还是码表登记不全
                    bads = [v for v, _n in (r.get('bad_values') or [])]
                    rv = real_values(cur, t, c, sp, pks, all_sim, bads)
                    lvl, fmtv, restv, codev = classify(r, rv)
                    r['level'] = lvl
                    r['fmt_vals'] = fmtv[:12]
                    r['real_vals'] = sorted(rv)[:12] if rv else []
                    r['real_n'] = len(rv) if rv is not None else -1
                    if not fmtv:
                        r['codes_sample'] = codev
                else:
                    r['level'] = 'OK'
                r.pop('_codes', None)
                r.pop('_names', None)
                results.append(r)
            except Exception as e:
                results.append({'table': t, 'column': c, 'form': sp['form'] or '(未标注)',
                                'code_names': ','.join(sorted(sp['code_names'])[:3]),
                                'src': '+'.join(sp['src']), 'error': str(e)[:160]})
        print(f"  [{i}/{len(work)}] {t:<46} 字段{len(cols):>3}  "
              f"{'全表仿真' if all_sim else '主键圈定(%d)' % len(pks)}")

    # ---------------- 汇总 ----------------
    ok = [r for r in results if 'error' not in r and r['n_bad'] == 0]
    bad = [r for r in results if 'error' not in r and r['n_bad'] > 0]
    err = [r for r in results if 'error' in r]
    print()
    print("=" * 78)
    print(f"字段总数 {len(results)}｜合规 {len(ok)}｜不合规 {len(bad)}｜查询失败 {len(err)}")
    if bad:
        from collections import Counter as _C
        print()
        print("按性质分级：")
        for lv, n in _C(r['level'] for r in bad).most_common():
            print(f"   {lv:<22} {n:>4} 个字段")
        print()
        print("不合规字段（按性质、越码行数降序）：")
        order = {'A 语义错误': 0, 'E 部分越码': 1, 'D 无对照(真实为空)': 2,
                 'D 无对照(全表仿真)': 3, 'B 格式差异': 4, 'C 码表不全(真实同值)': 5}
        for r in sorted(bad, key=lambda x: (order.get(x['level'], 9), -x['n_bad_rows']))[:70]:
            vals = ', '.join(f"{v}(x{n})" for v, n in r['bad_values'][:4])
            print(f"  [{r['level']:<20}] {r['table']}.{r['column']:<32} "
                  f"{r['n_bad_rows']:>8,}行  {vals}")
    if err:
        print()
        print("查询失败：")
        for r in err[:20]:
            print(f"  {r['table']}.{r['column']}  {r.get('error')}")

    if args.json_out:
        json.dump({'summary': {'total': len(results), 'ok': len(ok), 'bad': len(bad),
                               'err': len(err)},
                   'results': results},
                  open(args.json_out, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        print(f"\n明细已写 {args.json_out}")

    if args.out:
        write_xlsx(args.out, man, results, ok, bad, err)
        print(f"报告已写 {args.out}")

    conn.close()
    return 0


# ---------------------------------------------------------------- 输出
def write_xlsx(path, man, results, ok, bad, err):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    HF = PatternFill('solid', fgColor='1F4E79')
    HFONT = Font(color='FFFFFF', bold=True, size=10)
    RED = Font(color='C00000', bold=True)
    GREY = Font(color='808080')
    BD = Border(*[Side(style='thin', color='BFBFBF')] * 4)
    CTR = Alignment(horizontal='center', vertical='center')

    def hdr(ws, cols, widths):
        ws.append(cols)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for c in range(1, len(cols) + 1):
            cell = ws.cell(1, c)
            cell.fill = HF
            cell.font = HFONT
            cell.alignment = CTR
            cell.border = BD
        ws.freeze_panes = 'A2'

    # --- Sheet1 汇总 ---
    ws = wb.active
    ws.title = '检查结论'
    bad_tabs = OrderedDict()
    for r in bad:
        bad_tabs.setdefault(r['table'], []).append(r)
    ws.append(['deshu5 仿真数据 · 码值合规检查'])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(['仿真批次', f"{man.get('tag')}  {man.get('ym_from')}~{man.get('ym_to')}  "
                          f"创建于 {man.get('created_at')}"])
    ws.append(['检查表数', len({r['table'] for r in results})])
    ws.append(['检查字段数', f'{len(results)}  （仅含码表中有码项定义的字段）'])
    ws.append(['合规字段', len(ok)])
    ws.append(['不合规字段', len(bad)])
    ws.append(['查询失败字段', len(err)])
    ws.append(['不合规涉及表数', len(bad_tabs)])

    tot_rows = sum(r['total'] for r in results if 'error' not in r)
    bad_rows = sum(r['n_bad_rows'] for r in bad)
    ws.append([])
    ws.append(['总校验行数', f'{tot_rows:,}'])
    ws.append(['越码行数', f'{bad_rows:,}'])
    rate = (bad_rows / tot_rows * 100) if tot_rows else 0
    ws.append(['越码行占比', f'{rate:.4f}%'])

    ws.append([])
    ws.append(['违规分级', '字段数', '说明'])
    for c in range(1, 4):
        ws.cell(ws.max_row, c).fill = HF
        ws.cell(ws.max_row, c).font = HFONT
    LEVEL_DESC = {
        'A 语义错误': '仿真取值不在码表，真实数据亦无此值 —— 需修生成逻辑',
        'E 部分越码': '部分取值真实数据也有、部分独有 —— 需逐值确认',
        'D 无对照(真实为空)': '真实该字段全空，只能依码表判定 —— 建议按码表修正',
        'D 无对照(全表仿真)': '该表无真实行可对照',
        'B 格式差异': '码表登记形态与实际不一致（前导零/短码），需与业务确认真实存储形态',
        'C 码表不全(真实同值)': '真实数据同样存在这些值 —— 码表登记不全，非仿真缺陷',
    }
    from collections import Counter as _C
    for lv, n in _C(r['level'] for r in bad).most_common():
        ws.append([lv, n, LEVEL_DESC.get(lv, '')])
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
        if row[0].value:
            row[0].font = Font(bold=True)
            row[0].alignment = Alignment(horizontal='right', vertical='center')
    for i, w in enumerate([20, 62, 78], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # --- Sheet2 不合规明细 ---
    ws2 = wb.create_sheet('不合规明细')
    hdr(ws2, ['序号', '违规分级', '表名', '字段名', '存储形态', '码表', '码表来源',
              '仿真越码取值', '越码值种数', '越码行数', '字段总行数',
              '真实数据取值', '码表合法码值样例'],
        [5, 20, 42, 30, 10, 22, 22, 44, 10, 12, 12, 44, 34])
    order = {'A 语义错误': 0, 'E 部分越码': 1, 'D 无对照(真实为空)': 2,
             'D 无对照(全表仿真)': 3, 'B 格式差异': 4, 'C 码表不全(真实同值)': 5}
    for i, r in enumerate(sorted(bad, key=lambda x: (order.get(x['level'], 9), -x['n_bad_rows'])), 1):
        vals = ' | '.join(f'{v} (x{n})' for v, n in r['bad_values'][:8])
        realm = ' | '.join(str(x) for x in (r.get('real_vals') or []))
        ws2.append([i, r['level'], r['table'], r['column'], r['form'], r['code_names'],
                    r['src'], vals, r['n_bad'], r['n_bad_rows'], r['total'],
                    realm or ('—' if r.get('real_n') == 0 else '(不可对照)'),
                    ', '.join(r['codes_sample'][:10])])
        rr = ws2.max_row
        for c in range(1, 14):
            ws2.cell(rr, c).border = BD
        lv = r['level']
        if lv.startswith('A'):
            ws2.cell(rr, 2).font = RED
            ws2.cell(rr, 8).font = RED
        elif lv.startswith('B'):
            ws2.cell(rr, 2).font = Font(color='BF8F00', bold=True)
        elif lv.startswith('C'):
            ws2.cell(rr, 2).font = GREY
            ws2.cell(rr, 8).font = GREY
        else:
            ws2.cell(rr, 2).font = Font(color='BF8F00')
    ws2.auto_filter.ref = f"A1:M{ws2.max_row}"

    # --- Sheet3 全部字段 ---
    ws3 = wb.create_sheet('全部字段')
    hdr(ws3, ['表名', '字段名', '存储形态', '码表', '来源', '码项数(码/名)',
              '去重取值数', '空值行', '总行数', '越码值种数', '越码行数', '分级', '圈定方式'],
        [42, 30, 10, 22, 22, 14, 11, 10, 12, 11, 12, 20, 12])
    rows = sorted([r for r in results if 'error' not in r],
                  key=lambda x: (order.get(x.get('level'), -1), -x['n_bad_rows'],
                                 x['table'], x['column']))
    for r in rows:
        lv = r.get('level', 'OK') if r['n_bad'] else '合规'
        ws3.append([r['table'], r['column'], r['form'], r['code_names'], r['src'],
                    f"{r['n_code_items']}/{r['n_name_items']}", r['n_distinct'],
                    r['nulls'], r['total'], r['n_bad'], r['n_bad_rows'], lv,
                    '全表仿真' if r.get('all_sim') else '主键圈定'])
        rr = ws3.max_row
        for c in range(1, 14):
            ws3.cell(rr, c).border = BD
        if lv == '合规':
            ws3.cell(rr, 12).font = Font(color='2E7D32', bold=True)
        elif lv.startswith('A'):
            ws3.cell(rr, 12).font = RED
        elif lv.startswith('C'):
            ws3.cell(rr, 12).font = GREY
        else:
            ws3.cell(rr, 12).font = Font(color='BF8F00')
    ws3.auto_filter.ref = f"A1:M{ws3.max_row}"

    # --- Sheet4 查询失败 ---
    if err:
        ws4 = wb.create_sheet('查询失败')
        hdr(ws4, ['表名', '字段名', '存储形态', '码表', '错误'], [42, 30, 10, 22, 70])
        for r in err:
            ws4.append([r['table'], r['column'], r['form'], r.get('code_names', ''),
                        r.get('error', '')])

    wb.save(path)


if __name__ == '__main__':
    sys.exit(main())
