#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""growth_report.py — 生成「仿真前后条目数增长报告」xlsx。

对比两份全库行数快照（snapshot_counts.py 产出），并叠加本轮台账（manifest）信息，
得出：每张表基线行数 / 仿真写入 / 仿真后行数 / 增长倍数 / 所属域，以及覆盖率汇总。

用法：
    python growth_report.py --base before.json --after after.json \
        --tag 模拟 --out 仿真条目数增长报告.xlsx
"""
import argparse, glob, json, os, sys
from collections import OrderedDict

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("需要 openpyxl：pip install openpyxl")


# ---------- 分层判定 ----------
# 这两张是权威参考维表：所有外键都指向它们，真实数据本身就是口径来源，
# 刻意不仿真，单独成组以免在"主数据·营销域"里看起来像漏覆盖。
REF_TABLES = {'dim_cst_mgt_org', 'dim_org_region_map'}

LAYERS = [
    ('主数据·营销域', ('dim_cst_',)),
    ('业务数据·营销域', ('dwd_cst_',)),
    ('设备台账·PMS', ('dim_equ_', 'dim_ast_')),
    ('电网模型·D5000', ('dim_grid_', 'dwd_grid_')),
    ('统计汇总', ('ads_',)),
]


def layer_of(t):
    if t in REF_TABLES:
        return '参考维表（不仿真）'
    for name, pres in LAYERS:
        if t.startswith(pres):
            return name
    return '其他'


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--base', required=True, help='仿真前全库行数快照 json')
    ap.add_argument('--after', required=True, help='仿真后全库行数快照 json')
    ap.add_argument('--tag', default='模拟', help='本轮仿真的 tag，用于读台账')
    ap.add_argument('--manifest', default='', help='指定台账文件（默认取 manifest/ 下最新）')
    ap.add_argument('--out', default='仿真条目数增长报告.xlsx')
    ap.add_argument('--title', default='deshu5 仿真数据条目数增长报告')
    a = ap.parse_args()

    base = json.load(open(a.base, encoding='utf-8'))
    after = json.load(open(a.after, encoding='utf-8'))
    tables = sorted(set(base) | set(after))

    # 台账：哪些表引擎写出过行（rows=写入行数）
    mf = a.manifest
    if not mf:
        cands = sorted(glob.glob(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                              'manifest', f'*__{a.tag}__*.json')))
        mf = cands[-1] if cands else ''
    man = {}
    if mf and os.path.exists(mf):
        man = {k: v for k, v in json.load(open(mf, encoding='utf-8'))['tables'].items()}

    wb = openpyxl.Workbook()
    HDR = PatternFill('solid', fgColor='1F4E79')
    HF = Font(color='FFFFFF', bold=True, size=10)
    BOLD = Font(bold=True)
    RED = Font(color='C00000', bold=True)
    thin = Side(style='thin', color='BFBFBF')
    BD = Border(left=thin, right=thin, top=thin, bottom=thin)
    CTR = Alignment(horizontal='center', vertical='center')

    # ================= Sheet 1 汇总 =================
    ws = wb.active
    ws.title = '汇总'
    b_tot, a_tot = sum(base.values()), sum(after.values())
    n_all = len(tables)
    n_written = sum(1 for t in tables if (after.get(t, 0) - base.get(t, 0)) > 0)
    n_support = len(man)
    n_untouched = n_all - n_written

    ws.append([a.title])
    ws['A1'].font = Font(bold=True, size=14, color='1F4E79')
    ws.append([])
    rows = [
        ('快照基准', f"{a.base}"),
        ('快照结果', f"{a.after}"),
        ('本轮标记 tag', a.tag),
        ('台账文件', os.path.basename(mf) if mf else '（未找到）'),
        ('', ''),
        ('库内表总数', n_all),
        ('仿真前总行数', f"{b_tot:,}"),
        ('仿真后总行数', f"{a_tot:,}"),
        ('本轮新增行数', f"{a_tot - b_tot:,}"),
        ('整体增长倍数', f"{a_tot / b_tot:.2f} ×" if b_tot else '—'),
        ('', ''),
        ('本轮有写入的表数', n_written),
        ('本轮无写入的表数', n_untouched),
        ('台账登记（引擎支持）的表数', n_support),
        ('覆盖率（有写入/总表数）', f"{n_written / n_all * 100:.1f}%"),
    ]
    for k, v in rows:
        ws.append([k, v])
        if k:
            ws.cell(ws.max_row, 1).font = BOLD
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 46

    # 分层统计
    ws.append([])
    ws.append(['分层', '表数', '本轮有写入表数', '新增行数', '仿真后行数'])
    hr = ws.max_row
    for c in range(1, 6):
        cell = ws.cell(hr, c)
        cell.fill, cell.font, cell.alignment, cell.border = HDR, HF, CTR, BD
    for lname, _ in LAYERS + [('参考维表（不仿真）', ()), ('其他', ())]:
        ts = [t for t in tables if layer_of(t) == lname]
        if not ts:
            continue
        ws.append([lname, len(ts),
                   sum(1 for t in ts if after.get(t, 0) - base.get(t, 0) > 0),
                   sum(after.get(t, 0) - base.get(t, 0) for t in ts),
                   sum(after.get(t, 0) for t in ts)])
        for c in range(1, 6):
            ws.cell(ws.max_row, c).border = BD
    ws.append(['合计', n_all, n_written, a_tot - b_tot, a_tot])
    for c in range(1, 6):
        ws.cell(ws.max_row, c).font = BOLD

    # ================= Sheet 2 逐表明细 =================
    ws2 = wb.create_sheet('逐表明细')
    heads = ['序号', '分层', '表名', '仿真前', '本轮写入', '仿真后', '增长倍数', '台账登记', '主键']
    ws2.append(heads)
    for c in range(1, len(heads) + 1):
        cell = ws2.cell(1, c)
        cell.fill, cell.font, cell.alignment, cell.border = HDR, HF, CTR, BD
    order = OrderedDict((l, []) for l, _ in LAYERS + [('参考维表（不仿真）', ()), ('其他', ())])
    for t in tables:
        order[layer_of(t)].append(t)
    i = 0
    for lname, ts in order.items():
        for t in sorted(ts, key=lambda x: -(after.get(x, 0) - base.get(x, 0))):
            i += 1
            b, af = base.get(t, 0), after.get(t, 0)
            d = af - b
            m = man.get(t)
            ws2.append([i, lname, t, b, d, af,
                        round(af / b, 2) if b else ('新增' if af else '—'),
                        f"{m['rows']:,}" if m else '',
                        m['pk'] if m else ''])
            r = ws2.max_row
            for c in range(1, len(heads) + 1):
                ws2.cell(r, c).border = BD
            ws2.cell(r, 3).font = BOLD if d else Font(color='808080')
            if d:
                ws2.cell(r, 5).font = Font(color='C00000', bold=True)
    ws2.freeze_panes = 'A2'
    widths = [6, 18, 46, 12, 12, 12, 10, 12, 26]
    for idx, w in enumerate(widths, 1):
        ws2.column_dimensions[get_column_letter(idx)].width = w
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{ws2.max_row}"

    # ================= Sheet 3 未写入表 =================
    ws3 = wb.create_sheet('未写入表')
    ws3.append(['序号', '分层', '表名', '仿真前行数', '说明'])
    for c in range(1, 6):
        cell = ws3.cell(1, c)
        cell.fill, cell.font, cell.alignment, cell.border = HDR, HF, CTR, BD
    k = 0
    for t in tables:
        if after.get(t, 0) - base.get(t, 0) > 0:
            continue
        k += 1
        if t in man:
            why = '引擎支持，但本批未产生（无对应业务条件）'
        elif layer_of(t) == '参考维表（不仿真）':
            why = '权威参考维表：真实数据即口径来源，刻意不仿真'
        elif t.startswith('ads_'):
            why = '按口径决策不仿真：汇总层/外部统计（年鉴、城建等），非营销仿真范围'
        else:
            why = '引擎未覆盖'
        ws3.append([k, layer_of(t), t, base.get(t, 0), why])
        r = ws3.max_row
        for c in range(1, 6):
            ws3.cell(r, c).border = BD
        if why == '引擎未覆盖':
            ws3.cell(r, 5).font = RED
        elif t.startswith('ads_'):
            ws3.cell(r, 5).font = Font(color='808080')
    if k == 0:
        ws3.append(['—', '—', '（无）', '', '全部表均已写入'])
    for idx, w in enumerate([6, 18, 46, 12, 46], 1):
        ws3.column_dimensions[get_column_letter(idx)].width = w

    wb.save(a.out)
    print(f"[ok] 报告已生成：{a.out}")
    print(f"     表数 {n_all}｜有写入 {n_written}｜新增 {a_tot - b_tot:,} 行｜"
          f"{b_tot:,} → {a_tot:,}")


if __name__ == '__main__':
    main()
